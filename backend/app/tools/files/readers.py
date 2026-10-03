"""Text extraction per document type. Runs in a worker thread: parsing is CPU-bound."""

import asyncio
from collections.abc import Callable
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.core.errors import ToolInputError


def _pdf(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _docx(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + cells)


def _plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


_READERS: dict[str, Callable[[Path], str]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".csv": _plain,
    ".txt": _plain,
    ".json": _plain,
    ".md": _plain,
}


async def extract_text(path: Path) -> str:
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        raise ToolInputError(
            f"Unsupported file type '{path.suffix}'. Supported: {', '.join(_READERS)}"
        )
    return await asyncio.to_thread(_read_existing, path, reader)


def _read_existing(path: Path, reader: Callable[[Path], str]) -> str:
    if not path.is_file():
        raise ToolInputError(f"File '{path.name}' does not exist in the workspace")
    return reader(path)
